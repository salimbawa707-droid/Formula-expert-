#!/usr/bin/env python3
"""Export source-row crosswalk and reconcile 212 worksheet rows vs 210 data points.

The canonical workbook is authoritative for observed data records. The Round-6
parser counted 212 nonempty worksheet rows because it includes two header rows;
rows 3 onward contain 210 complete numeric flow points. This explains the
structural difference, but does not prove the original author's intent.
"""
from __future__ import annotations
import csv
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "evidence_vault/datasets/zenodo_18961871/Supplementary_SolutionFlowCurves.xlsx"
OUT = ROOT / "evidence_vault/forensic_reports"
CSV_OUT = OUT / "PHASE4B_FLOW_212_VS_210_SOURCE_ROW_CROSSWALK.csv"
JSON_OUT = OUT / "PHASE4B_FLOW_212_VS_210_SOURCE_ROW_CROSSWALK.json"
SHEET = "SolutionFlowCurves"
HEADER_ROWS = (1, 2)
HEADERS = ["filename", "sheet", "source_row", "source_record_id", "curve_label_raw",
           "point_value_raw", "numeric_point", "numeric_measurement_cells",
           "complete_numeric_point", "classification", "nonempty_cells"]

def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)

def safe(value):
    return "" if value is None else str(value).replace("\r", "\\r").replace("\n", "\\n")

def nonempty(values):
    return any(v is not None and str(v).strip() for v in values)

def main() -> int:
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Canonical source workbook missing: {SOURCE}")
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    workbook = load_workbook(SOURCE, data_only=False, read_only=False)
    if SHEET not in workbook.sheetnames:
        raise AssertionError(f"Expected worksheet {SHEET!r}; found {workbook.sheetnames!r}")
    ws = workbook[SHEET]
    header_nonempty = sum(nonempty([ws.cell(r,c).value for c in range(1,9)]) for r in HEADER_ROWS)
    rows, curve_labels = [], set()
    for r in range(3, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, 9)]
        if not nonempty(values):
            continue
        label, point, *measurements = values
        numeric_point = is_number(point)
        numeric_cells = [f"{ws.cell(r,c).coordinate}={safe(ws.cell(r,c).value)}"
                         for c in range(2, 9) if is_number(ws.cell(r,c).value)]
        complete = numeric_point and all(is_number(v) for v in measurements)
        if label is not None and str(label).strip():
            curve_labels.add(str(label).strip())
        rows.append({
            "filename": SOURCE.name, "sheet": SHEET, "source_row": r,
            "source_record_id": f"{SHEET}!A{r}:H{r}", "curve_label_raw": safe(label),
            "point_value_raw": safe(point), "numeric_point": str(numeric_point).lower(),
            "numeric_measurement_cells": ";".join(numeric_cells),
            "complete_numeric_point": str(complete).lower(),
            "classification": "COMPLETE_NUMERIC_FLOW_POINT" if complete else "NONEMPTY_ROW_NOT_COMPLETE_NUMERIC_POINT",
            "nonempty_cells": ";".join(f"{ws.cell(r,c).coordinate}={safe(values[c-1])}"
                                        for c in range(1, 9) if values[c-1] is not None),
        })
    max_row = ws.max_row
    workbook.close()
    nonempty_data_rows = len(rows)
    complete_points = sum(row["complete_numeric_point"] == "true" for row in rows)
    incomplete_rows = nonempty_data_rows - complete_points
    worksheet_nonempty_including_headers = nonempty_data_rows + header_nonempty
    OUT.mkdir(parents=True, exist_ok=True)
    with CSV_OUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)
    report = {
        "audit_id": "PHASE4B-FLOW-212-VS-210-SOURCE-ROW-CROSSWALK",
        "source": {"filename": SOURCE.name, "sheet": SHEET, "sha256": digest,
                   "source_range_scanned": f"A1:H{max_row}"},
        "counts": {
            "nonempty_header_rows": header_nonempty,
            "header_row_numbers": list(HEADER_ROWS),
            "observed_nonempty_data_rows": nonempty_data_rows,
            "complete_numeric_flow_points": complete_points,
            "nonempty_data_rows_not_complete_numeric_points": incomplete_rows,
            "observed_nonempty_worksheet_rows_including_headers": worksheet_nonempty_including_headers,
            "distinct_nonblank_curve_labels_seen": len(curve_labels)
        },
        "historical_claim_reconciliation": {
            "historical_claimed_nonempty_rows": 212,
            "historical_claimed_complete_points": 210,
            "worksheet_rows_including_headers_match_212": worksheet_nonempty_including_headers == 212,
            "data_points_match_210": complete_points == 210,
            "header_rows_account_for_difference": header_nonempty == 2 and worksheet_nonempty_including_headers - nonempty_data_rows == 2,
            "disposition": "STRUCTURALLY_RECONCILED_HEADER_ROWS_INCLUDED" if worksheet_nonempty_including_headers == 212 and complete_points == 210 else "NOT_REPRODUCED",
            "interpretation": "The 212 count is reproduced only when two nonempty header rows are included with 210 data rows. This is a source-structure explanation, not proof that the original author intentionally counted headers as observations."
        },
        "governance": {"phase4b": "OPEN", "phase5": "LOCKED"},
        "limitations": [
            "The 212 count includes two nonempty header rows and must not be described as 212 measured flow points.",
            "Structural row classification does not prove independent physical experiments.",
            "No missing or nonnumeric values are imputed.",
            "Original author intent remains unconfirmed without the original counting script or notebook."
        ],
        "csv_output": CSV_OUT.name
    }
    JSON_OUT.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
