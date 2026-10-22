#!/usr/bin/env python3
"""Round-19 raw workbook cell/formula/image/rheometer lineage audit."""
from __future__ import annotations
import csv, hashlib, json, re
from collections import Counter
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "evidence_vault" / "datasets" / "zenodo_18961871"
OUT = ROOT / "evidence_vault" / "forensic_reports"
FILES = ["Supplementary_SolutionFlowCurves.xlsx", "Supplementary_FilmThickness_v1.xlsx",
         "Supplementary_BubbleSizes_v1.xlsx", "Supplementary_RheometerData_v2.xlsx"]
CELL_CSV = OUT / "PHASE4B_ROUND19_CELL_FORMULA_LINEAGE.csv"
IMAGE_CSV = OUT / "PHASE4B_ROUND19_IMAGE_MEASUREMENT_LINEAGE.csv"
RHEO_CSV = OUT / "PHASE4B_ROUND19_RHEOMETER_TEST_POINT_LINEAGE.csv"
JSON_REPORT = OUT / "PHASE4B_ROUND19_RAW_WORKBOOK_LINEAGE.json"
MD_REPORT = OUT / "PHASE4B_ROUND19_RAW_WORKBOOK_LINEAGE.md"
CELL_HEADERS = ["filename","sheet","cell","row","column","cell_type","value_repr","formula","lineage_class"]
IMAGE_HEADERS = ["filename","sheet","row","image_group","image_id","number","length_px","film_thickness_um","record_status","source_cells"]
RHEO_HEADERS = ["filename","sheet","row","test","point_no","time","frequency","shear_rate","shear_stress","viscosity","record_status","source_cells"]

def safe(value: Any) -> str:
    if value is None: return ""
    return str(value).replace("\r","\\r").replace("\n","\\n")

def cell_class(filename: str, cell: Any) -> str:
    if cell.data_type == "f": return "FORMULA_DERIVED"
    name = filename.lower()
    if "bubblesizes" in name: return "IMAGE_MEASUREMENT_RAW_OR_DERIVED_INPUT"
    if "filmthickness" in name: return "IMAGE_LEVEL_RAW_OR_DERIVED_INPUT"
    if "rheometer" in name: return "RHEOMETER_TEST_DATA"
    if "solutionflow" in name: return "RHEOLOGY_FLOW_CURVE_DATA"
    return "OTHER"

def write_csv(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        w.writeheader(); w.writerows(rows)

def find_header_row(ws, required: set[str], scan_rows: int = 12):
    for r in range(1, min(ws.max_row, scan_rows)+1):
        mapping = {}
        for c in range(1, ws.max_column+1):
            val = ws.cell(r,c).value
            if isinstance(val,str) and val.strip():
                mapping[re.sub(r"\s+"," ",val.strip().lower())] = c
        if required.issubset(mapping): return r, mapping
    return None, {}

def image_lineage(filename: str, ws):
    if "filmthickness" not in filename.lower(): return []
    out = []
    groups = [(1,2,3,"Foam A"),(4,5,6,"Foam B"),(7,8,9,"Foam C"),(10,11,12,"Foam D")]
    for r in range(3, ws.max_row+1):
        for ic,nc,lc,label in groups:
            image, number, length = ws.cell(r,ic).value, ws.cell(r,nc).value, ws.cell(r,lc).value
            if image is None and number is None and length is None: continue
            if r < 3: continue
            derived = [f"{ws.cell(r,c).coordinate}={safe(ws.cell(r,c).value)}"
                       for c in range(13,ws.max_column+1) if ws.cell(r,c).value is not None]
            out.append({"filename":filename,"sheet":ws.title,"row":r,"image_group":label,
              "image_id":safe(image),"number":safe(number),"length_px":safe(length),"film_thickness_um":";".join(derived),
              "record_status":"TRACEABLE_IMAGE_ROW" if image is not None else "INCOMPLETE_NO_IMAGE_ID",
              "source_cells":f"{ws.cell(r,ic).coordinate};{ws.cell(r,nc).coordinate};{ws.cell(r,lc).coordinate}"})
    return out

def rheometer_lineage(filename: str, ws):
    if "rheometer" not in filename.lower() or ws.title.lower()=="contents": return []
    hr, mapping = find_header_row(ws, {"test","point no."})
    if hr is None: return []
    def col(*names):
        return next((mapping[n] for n in names if n in mapping), None)
    cols = {"test":col("test"),"point":col("point no."),"time":col("time"),
            "frequency":col("frequency"),"shear_rate":col("shear rate"),
            "shear_stress":col("shear stress"),"viscosity":col("viscosity","complex viscosity")}
    out=[]
    for r in range(hr+1,ws.max_row+1):
        vals={k:(ws.cell(r,c).value if c else None) for k,c in cols.items()}
        if all(v is None for v in vals.values()): continue
        coords=[ws.cell(r,c).coordinate for c in cols.values() if c]
        out.append({"filename":filename,"sheet":ws.title,"row":r,"test":safe(vals["test"]),
          "point_no":safe(vals["point"]),"time":safe(vals["time"]),"frequency":safe(vals["frequency"]),
          "shear_rate":safe(vals["shear_rate"]),"shear_stress":safe(vals["shear_stress"]),
          "viscosity":safe(vals["viscosity"]),
          "record_status":"CANDIDATE_TEST_POINT" if vals["point"] is not None else "NO_POINT_NUMBER",
          "source_cells":";".join(coords)})
    return out

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    cells=[]; images=[]; rheo=[]; workbook_stats=[]; missing=[]
    for filename in FILES:
        path=DATA/filename
        if not path.exists(): missing.append(filename); continue
        digest=hashlib.sha256(path.read_bytes()).hexdigest()
        wf=load_workbook(path,data_only=False,read_only=False)
        wv=load_workbook(path,data_only=True,read_only=False)
        fs={"filename":filename,"bytes":path.stat().st_size,"sha256":digest,"sheets":[]}
        for ws in wf.worksheets:
            vws=wv[ws.title]; nonempty=numeric=formulas=0
            for row in ws.iter_rows():
                for cell in row:
                    value=cell.value
                    if value is None: continue
                    nonempty+=1
                    if isinstance(value,(int,float)) and not isinstance(value,bool): numeric+=1
                    formula=safe(value) if cell.data_type=="f" else ""
                    if formula: formulas+=1
                    cached=vws[cell.coordinate].value if formula else value
                    cells.append({"filename":filename,"sheet":ws.title,"cell":cell.coordinate,
                      "row":cell.row,"column":cell.column,"cell_type":cell.data_type,
                      "value_repr":safe(cached),"formula":formula,"lineage_class":cell_class(filename,cell)})
            fs["sheets"].append({"sheet":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,
              "nonempty_cells":nonempty,"numeric_cells":numeric,"formula_cells":formulas})
            images.extend(image_lineage(filename,ws)); rheo.extend(rheometer_lineage(filename,ws))
        wf.close(); wv.close(); workbook_stats.append(fs)
    write_csv(CELL_CSV,CELL_HEADERS,cells); write_csv(IMAGE_CSV,IMAGE_HEADERS,images); write_csv(RHEO_CSV,RHEO_HEADERS,rheo)
    rcount=Counter((r["filename"],r["sheet"],r["record_status"]) for r in rheo)
    icount=Counter((r["filename"],r["image_group"],r["record_status"]) for r in images)
    totals={"nonempty_cell_records":len(cells),"formula_cell_records":sum(r["cell_type"]=="f" for r in cells),
      "image_lineage_records":len(images),"image_records_with_id":sum(r["record_status"]=="TRACEABLE_IMAGE_ROW" for r in images),
      "rheometer_candidate_rows":sum(r["record_status"]=="CANDIDATE_TEST_POINT" for r in rheo),
      "rheometer_rows_without_point_number":sum(r["record_status"]=="NO_POINT_NUMBER" for r in rheo)}
    report={"audit_id":"PHASE4B-ROUND19-RAW-WORKBOOK-LINEAGE","schema_version":1,
      "decision":"LINEAGE_EXTRACTED_AGGREGATE_STILL_REQUIRES_SEMANTIC_RECONCILIATION",
      "governance":{"phase4b":"OPEN","phase5":"LOCKED"},
      "historical_claim":{"claimed_total":27982,"independently_reproduced":False},
      "missing_workbooks":missing,"workbooks":workbook_stats,"totals":totals,
      "rheometer_rows_by_sheet_status":[{"filename":k[0],"sheet":k[1],"status":k[2],"rows":v} for k,v in sorted(rcount.items())],
      "image_rows_by_group_status":[{"filename":k[0],"group":k[1],"status":k[2],"rows":v} for k,v in sorted(icount.items())],
      "outputs":[CELL_CSV.name,IMAGE_CSV.name,RHEO_CSV.name],
      "guardrails":["A nonempty cell is not automatically a physical measurement.",
        "Formula-derived values are flagged separately; they are not added to raw-input counts.",
        "Image-row candidates are not assumed to be unique images or independent experiments.",
        "Rheometer rows are candidates only; Point No. duplicates and repeat grouping require explicit review.",
        "Formula cached values may be absent or stale; formulas and input cell coordinates are retained.",
        "No corrected total replaces 27,982 until category-specific lineage and counting semantics reconcile."]}
    JSON_REPORT.write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    lines=["# Phase 4B Round-19 — Raw Workbook Lineage Extraction","","**Decision:** LINEAGE EXTRACTED; historical aggregate still unresolved",
      "**Governance:** Phase 4B OPEN; Phase 5 LOCKED","","## Workbook integrity and structure","","| Workbook | Bytes | SHA-256 |","|---|---:|---|"]
    for w in workbook_stats: lines.append(f"| {w['filename']} | {w['bytes']} | {w['sha256']} |")
    lines += ["","## Extracted lineage inventory", "",f"- Non-empty cell records: **{totals['nonempty_cell_records']}**",
      f"- Formula cells flagged: **{totals['formula_cell_records']}**",
      f"- Film/image lineage rows: **{totals['image_lineage_records']}** (with image ID: {totals['image_records_with_id']})",
      f"- Rheometer candidate rows with Point No.: **{totals['rheometer_candidate_rows']}**",
      f"- Rheometer rows without Point No.: **{totals['rheometer_rows_without_point_number']}**",
      "","These are lineage inventory counts, not independent physical observations and not a replacement for 27,982.",
      "","## Required follow-up","","1. Review duplicate image IDs and image/number pairs; separate multiple features per image from distinct images.",
      "2. Inspect formula expressions and precedents; separate raw image measurements, calculated radii/thicknesses, and aggregate summaries.",
      "3. Group rheometer records only with a documented key and method-defined test/repeat semantics.",
      "4. Map each historical category component to explicit source records and document exclusions.",
      "5. Recalculate the historical total only after the mappings are independently reproducible.",
      "","## Guardrails","","- No physical observation count is inferred from row or numeric-cell count.",
      "- Formula-derived values are not automatically independent measurements.",
      "- No engine is promoted to LAB_VALIDATED.",
      "- **Phase 4B remains OPEN; Phase 5 remains LOCKED.**",
      "","Machine-readable detail is emitted as three CSV artifacts plus the JSON summary."]
    MD_REPORT.write_text("\n".join(lines)+"\n",encoding="utf-8")
    print(json.dumps({"decision":report["decision"],"totals":totals,"missing_workbooks":missing},indent=2))
    return 2 if missing else 0

if __name__=="__main__": raise SystemExit(main())
