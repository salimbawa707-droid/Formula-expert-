#!/usr/bin/env python3
"""Source-row crosswalk for the historical rheometer count claim of 739."""
from __future__ import annotations
import csv, hashlib, json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/"evidence_vault/datasets/zenodo_18961871/Supplementary_RheometerData_v2.xlsx"
OUT=ROOT/"evidence_vault/forensic_reports"
CSV_OUT=OUT/"PHASE4B_RHEOMETER_739_SOURCE_ROW_CROSSWALK.csv"
JSON_OUT=OUT/"PHASE4B_RHEOMETER_739_SOURCE_ROW_CROSSWALK.json"
HEADERS=["filename","sheet","source_row","test_label_raw","test_label_effective","point_no",
         "test_cell","point_no_cell","time","frequency","shear_rate","shear_stress","viscosity",
         "test_point_key_duplicate_count","counting_unit"]

def is_number(v): return isinstance(v,(int,float)) and not isinstance(v,bool)
def safe(v): return "" if v is None else str(v).replace("\r","\\r").replace("\n","\\n")

def main():
    if not SOURCE.is_file(): raise FileNotFoundError(f"Canonical rheometer workbook missing: {SOURCE}")
    digest=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    wb=load_workbook(SOURCE,data_only=False,read_only=False)
    records=[]; stats={}
    for ws in wb.worksheets:
        if ws.title.lower()=="contents": continue
        header=None; h={}
        for r in range(1,min(ws.max_row,12)+1):
            mapping={str(ws.cell(r,c).value).strip().lower():c for c in range(1,ws.max_column+1) if ws.cell(r,c).value is not None}
            if "test" in mapping and "point no." in mapping:
                header=r; h=mapping; break
        if header is None: continue
        last=""; sheet_records=[]; key_counts=Counter()
        for r in range(header+1,ws.max_row+1):
            raw_test=ws.cell(r,h["test"]).value
            if raw_test is not None and str(raw_test).strip(): last=str(raw_test).strip()
            point=ws.cell(r,h["point no."]).value
            if not is_number(point): continue
            def get(*names):
                col=next((h[n] for n in names if n in h),None)
                return safe(ws.cell(r,col).value) if col else ""
            key=(last,safe(point)); key_counts[key]+=1
            rec={"filename":SOURCE.name,"sheet":ws.title,"source_row":r,
                 "test_label_raw":safe(raw_test),"test_label_effective":last,"point_no":safe(point),
                 "test_cell":ws.cell(r,h["test"]).coordinate,"point_no_cell":ws.cell(r,h["point no."]).coordinate,
                 "time":get("time"),"frequency":get("frequency"),"shear_rate":get("shear rate"),
                 "shear_stress":get("shear stress"),"viscosity":get("viscosity","complex viscosity"),
                 "test_point_key_duplicate_count":"","counting_unit":"numeric Point No. source row; not automatically an independent experiment"}
            records.append(rec); sheet_records.append(rec)
        for rec in sheet_records:
            rec["test_point_key_duplicate_count"]=key_counts[(rec["test_label_effective"],rec["point_no"])]
        stats[ws.title]={"numeric_point_rows":len(sheet_records),
                         "distinct_forward_filled_test_groups":len({r["test_label_effective"] for r in sheet_records}),
                         "duplicate_test_point_keys":sum(n>1 for n in key_counts.values()),
                         "rows_with_blank_raw_test_label":sum(not r["test_label_raw"] for r in sheet_records)}
    wb.close()
    total=len(records); groups=sum(v["distinct_forward_filled_test_groups"] for v in stats.values())
    OUT.mkdir(parents=True,exist_ok=True)
    with CSV_OUT.open("w",newline="",encoding="utf-8") as f:
        w=csv.DictWriter(f,fieldnames=HEADERS); w.writeheader(); w.writerows(records)
    report={"audit_id":"PHASE4B-RHEOMETER-739-SOURCE-ROW-CROSSWALK",
            "source":{"filename":SOURCE.name,"sha256":digest},
            "counts":{"numeric_point_rows":total,"test_groups_across_sheets":groups,
                      "historical_claimed_rheometer_count":739,
                      "source_rows_minus_historical_claim":total-739,
                      "historical_739_reproduced":total==739},
            "worksheets":stats,
            "disposition":"NOT_REPRODUCED" if total!=739 else "COUNT_REPRODUCED_SEMANTICS_STILL_REQUIRE_REVIEW",
            "governance":{"phase4b":"OPEN","phase5":"LOCKED"},
            "limitations":["Numeric Point No. rows are source records, not proof of independent experiments.",
                           "Blank Test labels are forward-filled only for grouping and remain blank in the raw-label field.",
                           "The historical 739 claim has no demonstrated source-row mapping in this audit; no rows are dropped to force the claimed total."],
            "csv_output":CSV_OUT.name}
    JSON_OUT.write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(report,indent=2,ensure_ascii=False)); return 0

if __name__=="__main__": raise SystemExit(main())
