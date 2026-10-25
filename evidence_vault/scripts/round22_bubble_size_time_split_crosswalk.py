#!/usr/bin/env python3
"""Cell-level audit of bubble area measurements and formula-derived radii by time sheet."""
from __future__ import annotations
import csv, hashlib, json, re
from itertools import combinations
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "evidence_vault/datasets/zenodo_18961871/Supplementary_BubbleSizes_v1.xlsx"
OUT = ROOT / "evidence_vault/forensic_reports"
CSV_OUT = OUT / "PHASE4B_BUBBLE_20_30MIN_CELL_CROSSWALK.csv"
JSON_OUT = OUT / "PHASE4B_BUBBLE_20_30MIN_CELL_CROSSWALK.json"
TIME_SHEETS = {"20min": 10000, "30min": 9509}
HEADERS = ["filename","sheet","source_row","raw_area_cell","raw_area_value","derived_radius_cell",
           "derived_radius_formula","formula_precedents","raw_radius_pair_status","counting_unit"]

def is_number(v):
    return isinstance(v, (int,float)) and not isinstance(v,bool)

def safe(v):
    return "" if v is None else str(v).replace("\r","\\r").replace("\n","\\n")

def precedents(formula):
    return re.findall(r"(?<![A-Z0-9_])\$?([A-Z]{1,3})\$?(\d+)", formula.upper())

def main():
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Canonical bubble workbook missing: {SOURCE}")
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    wf = load_workbook(SOURCE, data_only=False, read_only=False)
    wv = load_workbook(SOURCE, data_only=True, read_only=False)
    missing = set(TIME_SHEETS) - set(wf.sheetnames)
    if missing:
        raise AssertionError(f"Expected time sheets {sorted(missing)}; found {wf.sheetnames}")
    records, by_sheet = [], {}
    for sheet, historical_count in TIME_SHEETS.items():
        ws, vs = wf[sheet], wv[sheet]
        raw_count = formula_count = matched = 0
        column_counts = {col: 0 for col in "ABCD"}
        for row in range(3, ws.max_row + 1):
            for raw_col, radius_col in zip("ABCD", "EFGH"):
                raw = ws[f"{raw_col}{row}"]
                formula = ws[f"{radius_col}{row}"]
                if not is_number(raw.value):
                    continue
                raw_count += 1
                column_counts[raw_col] += 1
                refs = precedents(str(formula.value)) if formula.data_type == "f" else []
                matched_here = formula.data_type == "f" and (raw_col, str(row)) in refs
                if formula.data_type == "f":
                    formula_count += 1
                if matched_here:
                    matched += 1
                records.append({
                    "filename": SOURCE.name, "sheet": sheet, "source_row": row,
                    "raw_area_cell": raw.coordinate, "raw_area_value": safe(raw.value),
                    "derived_radius_cell": formula.coordinate if formula.data_type == "f" else "",
                    "derived_radius_formula": safe(formula.value) if formula.data_type == "f" else "",
                    "formula_precedents": ";".join(f"{c}{r}" for c,r in refs),
                    "raw_radius_pair_status": "MATCHED_FORMULA_DEPENDENCY" if matched_here else "UNRESOLVED",
                    "counting_unit": "one raw area feature; derived radius is not an additional observation"
                })
        by_sheet[sheet] = {
            "observed_raw_numeric_area_cells": raw_count,
            "raw_area_numeric_cells_by_column": column_counts,
            "formula_cells_in_radius_columns_on_rows_scanned": formula_count,
            "matched_raw_to_radius_formula_pairs": matched,
            "historical_claimed_count": historical_count,
            "observed_minus_historical": raw_count - historical_count,
            "historical_count_reproduced": raw_count == historical_count,
        }
    # Search all eight source columns for arithmetic-only subset matches to the
    # historical 10,000 / 9,509 split. A match is a lead, never proof of intent.
    all_columns = {f"{sheet}!{col}": count
                   for sheet, stats in by_sheet.items()
                   for col, count in stats["raw_area_numeric_cells_by_column"].items()}
    keys = list(all_columns)
    subsets_10000, subsets_9509 = [], []
    for size in range(1, len(keys) + 1):
        for chosen in combinations(keys, size):
            subtotal = sum(all_columns[k] for k in chosen)
            if subtotal == 10000:
                subsets_10000.append(list(chosen))
            if subtotal == 9509:
                subsets_9509.append(list(chosen))
    paired_subset_hypotheses = []
    target_set = set(keys)
    for chosen_20 in subsets_10000:
        remainder = sorted(target_set - set(chosen_20))
        if sum(all_columns[k] for k in remainder) == 9509:
            paired_subset_hypotheses.append({
                "20min_claim_candidate_columns": chosen_20,
                "30min_claim_candidate_columns": remainder,
                "status": "ARITHMETIC_ONLY_NOT_EVIDENCE_OF_HISTORICAL_RULE"
            })
    wf.close(); wv.close()
    OUT.mkdir(parents=True, exist_ok=True)
    with CSV_OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=HEADERS); w.writeheader(); w.writerows(records)
    total = sum(v["observed_raw_numeric_area_cells"] for v in by_sheet.values())
    total_matched = sum(v["matched_raw_to_radius_formula_pairs"] for v in by_sheet.values())
    report = {
        "audit_id":"PHASE4B-BUBBLE-20-30MIN-CELL-CROSSWALK",
        "source":{"filename":SOURCE.name,"sha256":digest,"sheets":list(TIME_SHEETS)},
        "sheet_counts":by_sheet,
        "historical_split_column_hypothesis_search": {
            "source_column_numeric_counts": all_columns,
            "subset_count_matching_10000": len(subsets_10000),
            "subset_count_matching_9509": len(subsets_9509),
            "complementary_partition_count_matching_both": len(paired_subset_hypotheses),
            "complementary_partitions": paired_subset_hypotheses,
            "interpretation": "Column-subset arithmetic is diagnostic only. It cannot establish that the historical count used these columns or that the subset corresponds to a physical time point."
        },
        "aggregate":{"raw_numeric_area_cells":total,"matched_formula_pairs":total_matched,
                     "historical_aggregate_claim":19509,"aggregate_reproduced":total == 19509,
                     "formula_dependency_pairs_counted_as_independent_observations":False},
        "disposition":"AGGREGATE_MATCH_TIME_SPLIT_MISMATCH" if total == 19509 and not all(v["historical_count_reproduced"] for v in by_sheet.values()) else ("RECONCILED" if total == 19509 else "NOT_REPRODUCED"),
        "governance":{"phase4b":"OPEN","phase5":"LOCKED"},
        "limitations":[
            "A formula-derived radius is paired to its raw area input and is not counted as a second physical feature.",
            "Worksheet rows/cells are traceable records, not proof of independent experiments.",
            "Historical 20-minute and 30-minute split is tested against the source sheet counts; mismatch is not silently reallocated."
        ],
        "csv_output":CSV_OUT.name
    }
    JSON_OUT.write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(report,indent=2,ensure_ascii=False))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
