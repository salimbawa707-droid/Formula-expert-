#!/usr/bin/env python3
"""Reproduce the water-reference baseline check from Zenodo XLSX rows.

This is not an application-model prediction evaluator. The app returns a fixed
1 cP pure-water reference for an empty ingredient list, marked
reference_baseline. The 25 C reference is compared descriptively with
water-labelled measurements near 20 C; detergent samples are excluded.
"""
import argparse
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook

EXPECTED_SHA256 = "90ee8648151c2df423355650db1f82420317a53d4d3de1717ea9d4b34b61d590"
MODEL_COMMIT = "479bf76417de4c598734ae8bd3ccf77f3e086f7d"
MODEL_PATH = "src/utils/viscosityEngine.ts"
MODEL_BLOB_SHA = "fd8fe360a9cf923953df561366c615d953a84701"
REFERENCE_CP = 1.0


def metrics(predicted, observed):
    errors = [p - o for p, o in zip(predicted, observed)]
    return {
        "n": len(errors),
        "mae_cP": mean(abs(e) for e in errors) if errors else None,
        "rmse_cP": math.sqrt(mean(e * e for e in errors)) if errors else None,
        "bias_predicted_minus_measured_cP": mean(errors) if errors else None,
    }


def evaluate(workbook_path: Path):
    file_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    if file_hash != EXPECTED_SHA256:
        raise ValueError(f"Unexpected source workbook SHA-256: {file_hash}")
    sheet = load_workbook(workbook_path, data_only=True, read_only=True)["SolutionFlowCurves"]
    groups = defaultdict(list)
    current_label = None
    for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
        if row[0] is not None:
            current_label = str(row[0]).strip()
        if current_label and isinstance(row[3], (int, float)) and isinstance(row[5], (int, float)):
            groups[current_label].append({
                "source_row": row_number,
                "sample_label": current_label,
                "shear_rate_s-1": float(row[3]),
                "observed_viscosity_Pa_s": float(row[5]),
                "observed_viscosity_cP": float(row[5]) * 1000.0,
                "temperature_C": float(row[7]) if isinstance(row[7], (int, float)) else None,
            })
    water_groups = {k: v for k, v in groups.items() if k.startswith("Water_")}
    if set(water_groups) != {"Water_1", "Water_2", "Water_3"}:
        raise ValueError(f"Unexpected water sample labels: {sorted(water_groups)}")
    observations = [item for values in water_groups.values() for item in values]
    observed = [item["observed_viscosity_cP"] for item in observations]
    predicted = [REFERENCE_CP] * len(observed)
    per_curve = {}
    for name, rows in water_groups.items():
        curve_mean = mean(item["observed_viscosity_cP"] for item in rows)
        per_curve[name] = {
            "source_rows": [item["source_row"] for item in rows],
            "row_count": len(rows),
            "mean_observed_viscosity_cP": curve_mean,
            "temperature_min_C": min(item["temperature_C"] for item in rows),
            "temperature_max_C": max(item["temperature_C"] for item in rows),
            "shear_rate_min_s-1": min(item["shear_rate_s-1"] for item in rows),
            "shear_rate_max_s-1": max(item["shear_rate_s-1"] for item in rows),
            "metrics_vs_reference": metrics([REFERENCE_CP], [curve_mean]),
        }
    grouped_observed = [mean(item["observed_viscosity_cP"] for item in rows) for rows in water_groups.values()]
    return {
        "evaluation_id": "task2-water-reference-baseline-round1",
        "status": "DESCRIPTIVE_BASELINE_ONLY_NO_MODEL_VALIDATION",
        "source": {
            "dataset": "Zenodo record 18961871",
            "workbook": "Supplementary_SolutionFlowCurves.xlsx",
            "worksheet": "SolutionFlowCurves",
            "sha256": file_hash,
            "license": "CC BY 4.0",
            "unit_row": 2,
            "source_rows": [item["source_row"] for item in observations],
        },
        "model_reference": {
            "repository_commit": MODEL_COMMIT,
            "source_path": MODEL_PATH,
            "source_blob_sha": MODEL_BLOB_SHA,
            "application_entry_point": "ViscosityEngine.evaluateViscosity([])",
            "observed_source_contract": "createWaterBaselineReport() in the audited 479bf76417de4c598734ae8bd3ccf77f3e086f7d source snapshot returns viscosityCp=1 at 25 C; that snapshot mislabeled the reference measured_authoritative, corrected to reference_baseline by Task 2A.",
            "reference_value": REFERENCE_CP,
            "unit": "cP (mPa s)",
            "classification": "reference_baseline_not_formula_prediction",
        },
        "matching": {
            "variable": "dynamic viscosity",
            "conversion": "Pa s * 1000 = cP",
            "source_conditions": "Water_1 to Water_3 flow curves; measured temperature 19.99-20.00 C; shear rate is recorded per row",
            "reference_condition": "25 C fixed pure-water reference",
            "comparability": "descriptive_only_temperature_not_matched",
            "eligible_for_direct_formula_prediction_validation": 0,
            "formula_inputs_available_for_detergent_samples": False,
            "eligible_for_descriptive_reference_baseline": len(observations),
            "detergent_formula_rows_excluded": sum(len(v) for k, v in groups.items() if not k.startswith("Water_")),
        },
        "baseline_comparison": {
            "classification": "DESCRIPTIVE_BASELINE_COMPARISON_NOT_MODEL_VALIDATION",
            "pointwise_correlated_flow_curve_rows": metrics(predicted, observed),
            "curve_means_grouped_by_source_sample_label": metrics([REFERENCE_CP] * len(grouped_observed), grouped_observed),
        },
        "model_validation_metrics": None,
        "curves": per_curve,
        "lineage": observations,
        "validation": {
            "calibration_count": 0,
            "held_out_count": 0,
            "leakage_control": "No calibration or split was performed. Rows within each shear-rate curve were not treated as independent formula replicates; grouped curve summaries are reported separately.",
            "uncertainty": "No confidence interval reported: three source-labelled water curves, no per-measurement uncertainty or independent replicate provenance established.",
            "direct_model_validation": False,
            "foam_model_comparison": "not performed: ENG-005 registry implementation path is absent from audited repository commit and no matching executable foam kinetics model was located.",
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    result = evaluate(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": result["status"], "pointwise": result["baseline_comparison"]["pointwise_correlated_flow_curve_rows"], "curve_means": result["baseline_comparison"]["curve_means_grouped_by_source_sample_label"]}, indent=2))

if __name__ == "__main__":
    main()






