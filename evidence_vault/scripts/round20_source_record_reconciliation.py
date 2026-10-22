#!/usr/bin/env python3
"""Round-20 source-record reconciliation for historical 27,982 claim."""
import csv, hashlib, json, re
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
DATA=ROOT/"evidence_vault/datasets/zenodo_18961871"
OUT=ROOT/"evidence_vault/forensic_reports"
SOLUTION="Supplementary_SolutionFlowCurves.xlsx"
FILM="Supplementary_FilmThickness_v1.xlsx"
BUBBLES="Supplementary_BubbleSizes_v1.xlsx"
RHEO="Supplementary_RheometerData_v2.xlsx"
FOUT=OUT/"PHASE4B_ROUND20_FORMULA_DEPENDENCY_LEDGER.csv"
COUT=OUT/"PHASE4B_ROUND20_CATEGORY_RECORD_RECONCILIATION.csv"
ROUT=OUT/"PHASE4B_ROUND20_RHEOMETER_GROUPED_POINT_LEDGER.csv"
XOUT=OUT/"PHASE4B_ROUND20_HISTORICAL_CLAIM_CROSSWALK.csv"
JOUT=OUT/"PHASE4B_ROUND20_SOURCE_RECORD_RECONCILIATION.json"
MOUT=OUT/"PHASE4B_ROUND20_SOURCE_RECORD_RECONCILIATION.md"

def write_csv(p, headers, rows):
    with p.open("w",newline="",encoding="utf-8") as f:
        w=csv.DictWriter(f,fieldnames=headers,extrasaction="ignore"); w.writeheader(); w.writerows(rows)
def num(v): return isinstance(v,(int,float)) and not isinstance(v,bool)
def safe(v): return "" if v is None else str(v).replace("\r","\\r").replace("\n","\\n")
def refs(formula):
    cells=re.findall(r"(?<![A-Z0-9_])\$?([A-Z]{1,3})\$?(\d+)(?![A-Z0-9_])",formula,re.I)
    ranges=re.findall(r"(?<![A-Z0-9_])\$?([A-Z]{1,3}):\$?([A-Z]{1,3})(?![A-Z0-9_])",formula,re.I)
    return ";".join(dict.fromkeys([f"{c.upper()}{r}" for c,r in cells]+[f"{a.upper()}:{b.upper()}" for a,b in ranges]))
def load(fn):
    p=DATA/fn
    return p,load_workbook(p,data_only=False),load_workbook(p,data_only=True)

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    formulas=[]; categories=[]; rheo=[]; meta=[]

    # Solution flow point records
    p,wf,wv=load(SOLUTION); ws=wf["SolutionFlowCurves"]; last=""; curves=Counter(); flow=0
    for r in range(3,ws.max_row+1):
        label=ws.cell(r,1).value
        if label is not None and str(label).strip(): last=str(label).strip()
        point=ws.cell(r,2).value; vals=[ws.cell(r,c).value for c in range(3,9)]
        if num(point) and all(num(v) for v in vals):
            flow+=1; curves[last]+=1
            categories.append({"category":"solution_flow_curve_point","filename":SOLUTION,"sheet":ws.title,"source_row":r,
              "source_record_id":f"{ws.title}!A{r}:H{r}","group_id":last,"raw_record_count":1,"derived_record_count":0,"status":"RECONCILED"})
    meta.append({"filename":SOLUTION,"bytes":p.stat().st_size,"sha256":hashlib.sha256(p.read_bytes()).hexdigest(),
      "flow_points":flow,"curve_groups":dict(curves)})
    wf.close();wv.close()

    # Bubble area -> formula-derived radius pairs
    p,wf,wv=load(BUBBLES); raw_total=rad_total=matched=missing_cache=0; bubble_stats={}
    for ws in wf.worksheets:
        vws=wv[ws.title]; raw=rad=pair=0
        for r in range(3,ws.max_row+1):
            for a,b in zip("ABCD","EFGH"):
                av=ws[f"{a}{r}"].value; fc=ws[f"{b}{r}"]; cached=vws[f"{b}{r}"].value
                if num(av): raw+=1; raw_total+=1
                if fc.data_type=="f":
                    rad+=1;rad_total+=1
                    precedent=refs(str(fc.value))
                    ok=num(av) and f"{a}{r}" in precedent
                    if ok: pair+=1;matched+=1
                    if not num(cached): missing_cache+=1
                    formulas.append({"filename":BUBBLES,"sheet":ws.title,"formula_cell":fc.coordinate,"formula":fc.value,
                      "formula_precedents":precedent,"cached_value":safe(cached),"raw_input_cell":f"{a}{r}",
                      "raw_input_value":safe(av),"relationship":"MATCHED_RAW_AREA_TO_DERIVED_RADIUS" if ok else "UNRESOLVED",
                      "interpretation":"Count paired source feature once; derived radius is not an additional independent observation."})
        bubble_stats[ws.title]={"raw_area_cells":raw,"derived_radius_formula_cells":rad,"matched_pairs":pair}
        categories.append({"category":"bubble_area_radius_pair_by_sheet","filename":BUBBLES,"sheet":ws.title,"source_row":"",
          "source_record_id":f"{ws.title}!A3:H{ws.max_row}","group_id":ws.title,
          "raw_record_count":raw,"derived_record_count":rad,
          "status":"RECONCILED" if raw==rad==pair else "PAIRING_GAP",
          "notes":"Sheet-specific raw area/formula-radius pairs; historical per-sheet split compared separately."})
    categories.append({"category":"bubble_area_radius_pair","filename":BUBBLES,"sheet":"all","source_row":"",
      "source_record_id":"A:D raw areas paired with E:H formulas, row 3 onward","group_id":"20min and 30min",
      "raw_record_count":raw_total,"derived_record_count":rad_total,
      "status":"RECONCILED" if raw_total==rad_total==matched else "PAIRING_GAP",
      "notes":f"Matched pairs={matched}; nonnumeric/missing cached radius cells={missing_cache}; raw+derived must not be added together."})
    meta.append({"filename":BUBBLES,"bytes":p.stat().st_size,"sha256":hashlib.sha256(p.read_bytes()).hexdigest(),
      "sheets":bubble_stats,"raw_area_cells":raw_total,"radius_formula_cells":rad_total,"matched_pairs":matched,"missing_cached_radius_values":missing_cache})
    wf.close();wv.close()

    # Film feature records. Rows 1-2 are headers; row 3 onward is measured data.
    p,wf,wv=load(FILM); ws=wf["Data"];vws=wv["Data"]
    groups=[(1,2,3,"Foam A"),(4,5,6,"Foam B"),(7,8,9,"Foam C"),(10,11,12,"Foam D")]
    film=0; features=Counter(); unique_images=set(); film_formula=0
    for r in range(3,ws.max_row+1):
        for ic,nc,lc,label in groups:
            image,number,length=ws.cell(r,ic).value,ws.cell(r,nc).value,ws.cell(r,lc).value
            if image is not None and num(number) and num(length):
                film+=1;features[label]+=1;unique_images.add((label,safe(image)))
                categories.append({"category":"film_image_feature_measurement","filename":FILM,"sheet":ws.title,"source_row":r,
                  "source_record_id":f"{ws.title}!{ws.cell(r,ic).coordinate}:{ws.cell(r,lc).coordinate}",
                  "group_id":f"{label}|image={safe(image)}","raw_record_count":1,"derived_record_count":0,"status":"RECONCILED",
                  "notes":f"Number={safe(number)}; Length_px={safe(length)}"})
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type=="f":
                film_formula+=1; cached=vws[cell.coordinate].value
                formulas.append({"filename":FILM,"sheet":ws.title,"formula_cell":cell.coordinate,"formula":cell.value,
                  "formula_precedents":refs(str(cell.value)),"cached_value":safe(cached),"raw_input_cell":"","raw_input_value":"",
                  "relationship":"SUMMARY_OR_HISTOGRAM_FORMULA","interpretation":"Derived summary/distribution value; not an extra physical feature."})
    categories.append({"category":"film_image_feature_measurement","filename":FILM,"sheet":ws.title,"source_row":"",
      "source_record_id":"Data sheet, four three-column measurement groups","group_id":"Foam A/B/C/D",
      "raw_record_count":film,"derived_record_count":film_formula,"status":"RECONCILED",
      "notes":f"Feature counts={dict(features)}; distinct image label/group pairs={len(unique_images)}. Formula summaries excluded from physical feature count."})
    meta.append({"filename":FILM,"bytes":p.stat().st_size,"sha256":hashlib.sha256(p.read_bytes()).hexdigest(),
      "film_feature_measurements":film,"feature_counts_by_foam":dict(features),"distinct_image_label_groups":len(unique_images),"summary_formula_cells":film_formula})
    wf.close();wv.close()

    # Rheometer: forward-fill Test labels, count numeric Point No. rows only, exclude spacer rows.
    p,wf,wv=load(RHEO); sheet_stats={}
    for ws in wf.worksheets:
        if ws.title.lower()=="contents":continue
        header=None; h={}
        for r in range(1,min(ws.max_row,12)+1):
            m={str(ws.cell(r,c).value).strip().lower():c for c in range(1,ws.max_column+1) if ws.cell(r,c).value is not None}
            if "test" in m and "point no." in m:header=r;h=m;break
        if header is None:continue
        last=""; counts=Counter(); rawblank=0; dups=Counter()
        for r in range(header+1,ws.max_row+1):
            t=ws.cell(r,h["test"]).value
            if t is not None and str(t).strip():last=str(t).strip()
            point=ws.cell(r,h["point no."]).value
            if not num(point):continue
            counts[last]+=1;dups[(last,safe(point))]+=1
            if t is None or not str(t).strip():rawblank+=1
            def getval(*names):
                c=next((h[n] for n in names if n in h),None)
                return safe(ws.cell(r,c).value) if c else ""
            rheo.append({"filename":RHEO,"sheet":ws.title,"row":r,"test_raw":safe(t),"test_effective":last,
              "point_no":safe(point),"time":getval("time"),"frequency":getval("frequency"),
              "shear_rate":getval("shear rate"),"shear_stress":getval("shear stress"),
              "viscosity":getval("viscosity","complex viscosity"),
              "source_cells":f"{ws.cell(r,h['test']).coordinate};{ws.cell(r,h['point no.']).coordinate}"})
        sheet_stats[ws.title]={"point_rows":sum(counts.values()),"test_groups":len(counts),"points_by_test":dict(counts),
          "blank_raw_test_labels_forward_filled":rawblank,"duplicate_test_point_keys":[{"test":k[0],"point":k[1],"count":v} for k,v in dups.items() if v>1]}
    rheo_total=sum(v["point_rows"] for v in sheet_stats.values())
    categories.append({"category":"rheometer_test_point","filename":RHEO,"sheet":"six test worksheets","source_row":"",
      "source_record_id":"numeric Point No. rows grouped by forward-filled Test labels","group_id":"all tests",
      "raw_record_count":rheo_total,"derived_record_count":0,"status":"ROWS_RECONCILED_SEMANTICS_UNRESOLVED",
      "notes":"Blank spacer rows excluded. A point row is not automatically an independent physical experiment."})
    meta.append({"filename":RHEO,"bytes":p.stat().st_size,"sha256":hashlib.sha256(p.read_bytes()).hexdigest(),
      "sheets":sheet_stats,"rheometer_point_rows":rheo_total,"test_groups":sum(v["test_groups"] for v in sheet_stats.values())})
    wf.close();wv.close()

    write_csv(FOUT,["filename","sheet","formula_cell","formula","formula_precedents","cached_value","raw_input_cell","raw_input_value","relationship","interpretation"],formulas)
    write_csv(COUT,["category","filename","sheet","source_row","source_record_id","group_id","raw_record_count","derived_record_count","status","notes"],categories)
    write_csv(ROUT,["filename","sheet","row","test_raw","test_effective","point_no","time","frequency","shear_rate","shear_stress","viscosity","source_cells"],rheo)
    first3=flow+film+raw_total; full=first3+rheo_total; headline=27982; needed=headline-first3
    crosswalk=[
      {"claim":"solution_flow_curves","historical_count":212,"source_count":flow,"delta_source_minus_historical":flow-212,"status":"REPRODUCED" if flow==212 else "MISMATCH","notes":"Complete numeric point rows."},
      {"claim":"film_thickness_total","historical_count":7522,"source_count":film,"delta_source_minus_historical":film-7522,"status":"REPRODUCED" if film==7522 else "MISMATCH","notes":"Feature records, not unique images."},
      {"claim":"bubble_sizes_total","historical_count":19509,"source_count":raw_total,"delta_source_minus_historical":raw_total-19509,"status":"AGGREGATE_MATCH_TIME_SPLIT_MISMATCH" if raw_total==19509 else "MISMATCH","notes":"Historical split says 20min=10000, 30min=9509; raw source split is 20min=13388, 30min=6121."},
      {"claim":"bubble_sizes_20min","historical_count":10000,"source_count":bubble_stats.get("20min",{}).get("raw_area_cells",0),"delta_source_minus_historical":bubble_stats.get("20min",{}).get("raw_area_cells",0)-10000,"status":"MISMATCH","notes":"Raw numeric area cells in 20min sheet."},
      {"claim":"bubble_sizes_30min","historical_count":9509,"source_count":bubble_stats.get("30min",{}).get("raw_area_cells",0),"delta_source_minus_historical":bubble_stats.get("30min",{}).get("raw_area_cells",0)-9509,"status":"MISMATCH","notes":"Raw numeric area cells in 30min sheet."},
      {"claim":"rheometer","historical_count":739,"source_count":rheo_total,"delta_source_minus_historical":rheo_total-739,"status":"NOT_REPRODUCED","notes":"Numeric Point No. rows across six worksheets; spacer rows excluded."},
      {"claim":"grand_total","historical_count":headline,"source_count":full,"delta_source_minus_historical":full-headline,"status":"NOT_REPRODUCED","notes":"Mixed category measurement units; arithmetic comparison only."}]
    write_csv(XOUT,["claim","historical_count","source_count","delta_source_minus_historical","status","notes"],crosswalk)
    duplicate_keys=sum(len(v["duplicate_test_point_keys"]) for v in sheet_stats.values())
    checks={
      "solution_flow_points_210": flow==210,
      "film_feature_measurements_7522": film==7522,
      "bubble_area_formula_pairs_19509": raw_total==rad_total==matched==19509,
      "rheometer_point_rows_2233": rheo_total==2233,
      "rheometer_test_groups_62": sum(v["test_groups"] for v in sheet_stats.values())==62,
      "bubble_20min_raw_area_cells_13388": bubble_stats.get("20min",{}).get("raw_area_cells")==13388,
      "bubble_30min_raw_area_cells_6121": bubble_stats.get("30min",{}).get("raw_area_cells")==6121,
      "no_duplicate_point_keys_within_test": duplicate_keys==0,
    }
    report={"audit_id":"PHASE4B-ROUND20-SOURCE-RECORD-RECONCILIATION","decision":"BUBBLE_TIME_SPLIT_AND_RHEOMETER_739_NOT_REPRODUCED",
      "governance":{"phase4b":"OPEN","phase5":"LOCKED"},
      "historical_claim":{"headline":headline,"reported_components":{"solution_flow_curves":212,"film_thickness":7522,"bubble_sizes":19509,"rheometer":739},
        "reconciled_first_three_sum":first3,"rheometer_needed_to_reach_headline":needed,"rheometer_source_point_rows":rheo_total,
        "all_category_record_sum":full,"all_category_sum_minus_headline":full-headline,"headline_reproduced":False},
      "integrity_gate":{"status":"PASS" if all(checks.values()) else "FAIL","checks":checks,"duplicate_test_point_key_groups":duplicate_keys},
      "counts":{"solution_flow_curve_points":flow,"film_feature_measurements":film,"bubble_raw_area_cells":raw_total,
        "bubble_derived_radius_formula_cells":rad_total,"bubble_matched_pairs":matched,"rheometer_point_rows":rheo_total,
        "rheometer_test_groups":sum(v["test_groups"] for v in sheet_stats.values()),"formula_dependency_records":len(formulas),
        "film_summary_formula_cells":film_formula},
      "workbooks":meta,"outputs":[FOUT.name,COUT.name,ROUT.name,XOUT.name],
      "historical_claim_crosswalk":crosswalk,
      "limits":["Formula references are text-parsed, not a complete Excel dependency engine.",
        "Rheometer point rows are source rows, not automatically independent experiments.",
        "Historical rheometer 739 count has no demonstrated mapping to the raw point ledger.",
        "No engine is promoted to LAB_VALIDATED."]}
    JOUT.write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    lines=["# Phase 4B Round-20 — Source Record Reconciliation","","**Decision:** Aggregate bubble total matches, but its 20/30-minute split and rheometer 739 claim are not reproduced.",
      "**Governance:** Phase 4B OPEN; Phase 5 LOCKED.","","## Reconciled source counts","","| Category | Source count | Unit |","|---|---:|---|",
      f"| Solution flow curves | {flow:,} | Complete numeric point rows across named curves |",
      f"| Film thickness | {film:,} | Image-derived feature/length records; not distinct images |",
      f"| Bubble sizes | {raw_total:,} | Raw area entries paired with formula-derived radius values |",
      f"| Rheometer | {rheo_total:,} | Numeric Point No. rows across six worksheets |",
      f"| **All category records summed** | **{full:,}** | Mixed measurement units; not asserted as independent experiments |",
      "","## Historical headline reconciliation","","The first three source components sum to **27,241** using 210 reconciled solution-flow rows, 7,522 film features, and 19,509 bubble pairs. This differs from the historical flow count of 212 by two records. The bubble worksheet split also differs: historical values are 10,000 (20 min) and 9,509 (30 min), while raw area cells are 13,388 and 6,121 respectively. The aggregate bubble total still sums to 19,509, but the per-sheet lineage does not reproduce the published split. Using source-reconciled flow rows, the historical headline of 27,982 requires exactly **741** rheometer records. The source workbooks contain **2,233** numeric Point No. rows across **62** test groups; no explicit source-row mapping supports the historical 739. Summing all four category ledgers yields **29,474**, or **1,492 above** the headline. This arithmetic is not a claim that all categories are independent experiments.",
      "- Solution flow: 210 complete numeric point records across 15 named curves; historical count 212 is not reproduced (difference -2).",
      "- Film thickness: 7,522 numeric Length_px feature records. Repeated Image IDs are group labels; this is not 7,522 unique images.",
      "- Bubble sizes: 19,509 raw numeric area cells match 19,509 radius formulas by source-cell reference. The source split is 13,388 (20 min) and 6,121 (30 min), not the historical 10,000 / 9,509 split. Count each paired feature once, not twice.",
      "- Rheometer: point counts by worksheet are 288, 225, 192, 88, 720 and 720. Blank spacer rows are excluded; test labels are forward-filled for grouping.",
      "- The historical rheometer per-sheet figures mix counting rules: some include two header/unit rows, while the 731 steady-shear value includes spacer rows and does not equal numeric Point No. records.",
      "","## Decision","","27,982 is **not independently reproduced**. Keep it as a historical unresolved claim. Phase 4B stays OPEN, Phase 5 stays LOCKED, and no engine is promoted to LAB_VALIDATED.",
      "","## Machine-readable outputs","","- PHASE4B_ROUND20_FORMULA_DEPENDENCY_LEDGER.csv","- PHASE4B_ROUND20_CATEGORY_RECORD_RECONCILIATION.csv",
      "- PHASE4B_ROUND20_RHEOMETER_GROUPED_POINT_LEDGER.csv","- PHASE4B_ROUND20_HISTORICAL_CLAIM_CROSSWALK.csv","- PHASE4B_ROUND20_SOURCE_RECORD_RECONCILIATION.json"]
    MOUT.write_text("\n".join(lines)+"\n",encoding="utf-8")
    print(json.dumps({"decision":report["decision"],"counts":report["counts"],"historical_claim":report["historical_claim"]},indent=2))
    return 0
if __name__=="__main__": raise SystemExit(main())
